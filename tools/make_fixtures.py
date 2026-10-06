"""Generate the SlideLite test presentation collection.

    python3 tools/make_fixtures.py tests/fixtures

Requires python-pptx (dev only).  Every fixture is synthetic: no real-world
decks are committed to the repository.
"""

from __future__ import annotations

import io
import shutil
import struct
import sys
import zipfile
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures")
OUT.mkdir(parents=True, exist_ok=True)


def png(width: int, height: int, rgb=(200, 60, 60), alpha: bool = False) -> bytes:
    """Tiny dependency-free PNG encoder (a diagonal gradient)."""
    rows = []
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            t = (x + y) / max(1, width + height - 2)
            px = [int(c * (1 - t) + 255 * t) for c in rgb]
            if alpha:
                px.append(int(255 * (1 - x / max(1, width - 1))))
            row.extend(px)
        rows.append(bytes(row))
    raw = zlib.compress(b"".join(rows), 9)

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 6 if alpha else 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", raw) + chunk(b"IEND", b"")


def widescreen(prs: Presentation) -> Presentation:
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def save(prs: Presentation, name: str) -> Path:
    path = OUT / name
    prs.save(path)
    return path


# -- basic ----------------------------------------------------------------------


def simple_text() -> None:
    prs = widescreen(Presentation())
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "SlideLite Test Deck"
    slide.placeholders[1].text = "Simple text rendering"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Bullets and formatting"
    body = slide.placeholders[1].text_frame
    body.text = "First level bullet"
    for level, text in ((1, "Second level"), (2, "Third level"), (0, "Back to first")):
        p = body.add_paragraph()
        p.text = text
        p.level = level
    p = body.add_paragraph()
    for text, kw in (
        ("Bold ", {"bold": True}),
        ("italic ", {"italic": True}),
        ("underline ", {"underline": True}),
        ("red ", {"color": RGBColor(0xC0, 0x20, 0x20)}),
        ("big", {"size": Pt(36)}),
    ):
        run = p.add_run()
        run.text = text
        if "bold" in kw:
            run.font.bold = True
        if "italic" in kw:
            run.font.italic = True
        if "underline" in kw:
            run.font.underline = True
        if "color" in kw:
            run.font.color.rgb = kw["color"]
        if "size" in kw:
            run.font.size = kw["size"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = "Centered text box with a hyperlink"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = "example.com"
    run.hyperlink.address = "https://example.com/"
    save(prs, "simple-text.pptx")


def images() -> None:
    prs = widescreen(Presentation())
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(png(160, 90)), Inches(0.5), Inches(0.5), Inches(4))
    pic = slide.shapes.add_picture(
        io.BytesIO(png(120, 120, (40, 90, 200), alpha=True)), Inches(5), Inches(0.5), Inches(3)
    )
    pic.rotation = 15
    cropped = slide.shapes.add_picture(
        io.BytesIO(png(200, 100, (30, 160, 80))), Inches(9), Inches(0.5), Inches(3), Inches(3)
    )
    cropped.crop_left = 0.25
    cropped.crop_right = 0.25
    save(prs, "images.pptx")


def shapes() -> None:
    prs = widescreen(Presentation())
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    kinds = [
        MSO_SHAPE.RECTANGLE,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        MSO_SHAPE.OVAL,
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        MSO_SHAPE.RIGHT_ARROW,
        MSO_SHAPE.STAR_5_POINT,
        MSO_SHAPE.RECTANGULAR_CALLOUT,
        MSO_SHAPE.FLOWCHART_DECISION,
        MSO_SHAPE.HEXAGON,
        MSO_SHAPE.CHEVRON,
        MSO_SHAPE.HEART,
        MSO_SHAPE.CAN,
    ]
    for i, kind in enumerate(kinds):
        x = Inches(0.4 + (i % 6) * 2.15)
        y = Inches(0.6 + (i // 6) * 3.2)
        shape = slide.shapes.add_shape(kind, x, y, Inches(1.8), Inches(1.6))
        shape.text = kind.name.title().replace("_", " ")[:14]
        if i % 3 == 1:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0xF2, 0x8C, 0x28)
            shape.line.color.rgb = RGBColor(0x40, 0x40, 0x40)
            shape.line.width = Pt(3)
        if i == 5:
            shape.rotation = 20
    line = slide.shapes.add_connector(1, Inches(0.5), Inches(7), Inches(12.5), Inches(6.6))
    line.line.width = Pt(2)
    save(prs, "shapes.pptx")


def tables() -> None:
    prs = widescreen(Presentation())
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly results"
    rows, cols = 4, 4
    table = slide.shapes.add_table(rows, cols, Inches(1), Inches(1.8), Inches(11), Inches(3)).table
    data = [
        ["Region", "Q1", "Q2", "Q3"],
        ["North", "12", "15", "19"],
        ["South", "9", "11", "8"],
        ["Total", "21", "26", "27"],
    ]
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = data[r][c]
    table.cell(3, 0).merge(table.cell(3, 1))
    save(prs, "tables.pptx")


def multi_slide(count: int = 12, name: str = "multi-slide.pptx") -> None:
    prs = widescreen(Presentation())
    for i in range(count):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i + 1}"
        slide.placeholders[1].text = f"Content for slide {i + 1} of {count}"
    save(prs, name)


def aspect_ratios() -> None:
    prs = Presentation()  # python-pptx default is 4:3
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "4:3 presentation"
    save(prs, "aspect-4x3.pptx")
    prs = widescreen(Presentation())
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "16:9 presentation"
    save(prs, "aspect-16x9.pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(7560000), Emu(10692000)  # A4 portrait
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Custom portrait size"
    save(prs, "aspect-custom.pptx")


def _xml(fragment: str):
    """Parse a DrawingML fragment written with the a:/p: prefixes."""
    from pptx.oxml import parse_xml

    ns = (
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    )
    head, _, rest = (
        fragment.partition(" ")
        if fragment.index(" ") < fragment.index(">")
        else fragment.partition(">")
    )
    sep = " " if fragment.index(" ") < fragment.index(">") else ">"
    return parse_xml(f"{head}{ns}{sep}{rest}")


def text_formatting() -> None:
    prs = widescreen(Presentation())
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(6), Inches(6))
    tf = box.text_frame
    tf.word_wrap = True
    samples = [
        ("Liberation Serif 28pt", {"name": "Times New Roman", "size": Pt(28)}),
        ("Courier New 20pt", {"name": "Courier New", "size": Pt(20)}),
        ("Strikethrough", {"strike": True}),
        ("Superscript x2", {"baseline": "30000"}),
        ("Highlighted", {"highlight": True}),
        ("Letter spacing", {"spc": "600"}),
        ("ALL CAPS via cap", {"cap": "all"}),
    ]
    tf.text = "Text formatting"
    for text, opts in samples:
        p = tf.add_paragraph()
        run = p.add_run()
        run.text = text
        if "name" in opts:
            run.font.name = opts["name"]
        if "size" in opts:
            run.font.size = opts["size"]
        rpr = run._r.get_or_add_rPr()
        if opts.get("strike"):
            rpr.set("strike", "sngStrike")
        if opts.get("baseline"):
            rpr.set("baseline", opts["baseline"])
        if opts.get("spc"):
            rpr.set("spc", opts["spc"])
        if opts.get("cap"):
            rpr.set("cap", opts["cap"])
        if opts.get("highlight"):
            rpr.append(_xml('<a:highlight><a:srgbClr val="FFFF00"/></a:highlight>'))
    numbered = slide.shapes.add_textbox(Inches(7), Inches(0.4), Inches(5.5), Inches(3))
    ntf = numbered.text_frame
    ntf.text = "First numbered"
    for text in ("Second numbered", "Third numbered"):
        ntf.add_paragraph().text = text
    for p in ntf.paragraphs:
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", "342900")
        ppr.set("indent", "-342900")
        ppr.append(_xml('<a:buFont typeface="+mj-lt"/>'))
        ppr.append(_xml('<a:buAutoNum type="arabicPeriod"/>'))
        p.line_spacing = 1.5
    vertical = slide.shapes.add_textbox(Inches(7), Inches(4), Inches(1), Inches(3))
    vertical.text_frame.text = "Vertical text"
    vertical.text_frame._txBody.bodyPr.set("vert", "vert")
    rotated = slide.shapes.add_textbox(Inches(9), Inches(4.5), Inches(3), Inches(1))
    rotated.text_frame.text = "Rotated 30°"
    rotated.rotation = 30
    wing = slide.shapes.add_textbox(Inches(9), Inches(6), Inches(4), Inches(1))
    wing.text_frame.text = "Wingdings bullet"
    ppr = wing.text_frame.paragraphs[0]._p.get_or_add_pPr()
    ppr.set("marL", "285750")
    ppr.set("indent", "-285750")
    ppr.append(_xml('<a:buFont typeface="Wingdings" pitchFamily="2" charset="2"/>'))
    ppr.append(_xml('<a:buChar char="&#167;"/>'))
    save(prs, "text-formatting.pptx")


def fills() -> None:
    prs = widescreen(Presentation())
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.gradient()
    slide.background.fill.gradient_angle = 90
    stops = slide.background.fill.gradient_stops
    stops[0].color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    stops[1].color.rgb = RGBColor(0xC8, 0xD8, 0xF0)
    solid = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2)
    )
    solid.fill.solid()
    solid.fill.fore_color.rgb = RGBColor(0x20, 0x80, 0x40)
    grad = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4), Inches(0.5), Inches(3), Inches(2)
    )
    grad.fill.gradient()
    grad.fill.gradient_angle = 45
    patt = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.5), Inches(0.5), Inches(3), Inches(2))
    from pptx.enum.dml import MSO_PATTERN

    patt.fill.patterned()
    patt.fill.pattern = MSO_PATTERN.DIAGONAL_BRICK
    patt.fill.fore_color.rgb = RGBColor(0x80, 0x20, 0x20)
    patt.fill.back_color.rgb = RGBColor(0xFF, 0xF0, 0xE0)
    transparent = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.5), Inches(1.5), Inches(3), Inches(3)
    )
    transparent.fill.solid()
    transparent.fill.fore_color.rgb = RGBColor(0x30, 0x60, 0xF0)
    srgb = transparent.fill._xPr.find(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"
    )
    srgb.append(_xml('<a:alpha val="50000"/>'))
    dashed = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(5), Inches(3.5), Inches(3), Inches(2)
    )
    dashed.fill.background()
    dashed.line.color.rgb = RGBColor(0xC0, 0x30, 0x30)
    dashed.line.width = Pt(4)
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    dashed.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shadowed = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(9), Inches(3.5), Inches(3), Inches(2)
    )
    sppr = shadowed._element.spPr
    sppr.append(
        _xml(
            '<a:effectLst><a:outerShdw blurRad="76200" dist="38100" dir="2700000" algn="tl">'
            '<a:prstClr val="black"><a:alpha val="40000"/></a:prstClr></a:outerShdw></a:effectLst>'
        )
    )
    arrow = slide.shapes.add_connector(1, Inches(0.5), Inches(6.5), Inches(8), Inches(6.5))
    arrow.line.width = Pt(3)
    ln = arrow.line._get_or_add_ln()
    ln.append(_xml('<a:tailEnd type="triangle" w="med" len="med"/>'))
    save(prs, "fills.pptx")


def groups() -> None:
    prs = widescreen(Presentation())
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    a = outer.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    a.text = "A"
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.5), Inches(1), Inches(1), Inches(1)).text = "B"
    inner.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.5), Inches(2.2), Inches(2), Inches(0.8))
    flipped = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE, Inches(7), Inches(1), Inches(2), Inches(2)
    )
    flipped._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm").set(
        "flipH", "1"
    )
    save(prs, "groups.pptx")


# -- broken / hostile ------------------------------------------------------------------


def rewrite(src: Path, dst: str, edit) -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = edit(info.filename, zin.read(info))
            if data is not None:
                zout.writestr(info.filename, data)
    (OUT / dst).write_bytes(out.getvalue())


def broken() -> None:
    good = OUT / "multi-slide.pptx"
    data = good.read_bytes()
    (OUT / "corrupt-truncated.pptx").write_bytes(data[: len(data) // 2])
    (OUT / "corrupt-random.pptx").write_bytes(bytes((i * 37 + 11) % 256 for i in range(4096)))
    (OUT / "corrupt-empty.pptx").write_bytes(b"")
    rewrite(
        good, "corrupt-missing-slide.pptx", lambda n, d: None if n == "ppt/slides/slide2.xml" else d
    )
    rewrite(
        good,
        "corrupt-bad-xml.pptx",
        lambda n, d: d[: len(d) // 2] if n == "ppt/slides/slide3.xml" else d,
    )
    rewrite(
        good,
        "hostile-xxe.pptx",
        lambda n, d: (
            b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>'
            + d.split(b"?>", 1)[1]
            if n == "ppt/slides/slide1.xml"
            else d
        ),
    )

    def traversal(name: str, d: bytes):
        if name == "ppt/slides/_rels/slide1.xml.rels":
            return d.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../../../../../../etc/passwd"/></Relationships>',
            )
        return d

    rewrite(good, "hostile-traversal.pptx", traversal)
    bomb = io.BytesIO()
    with zipfile.ZipFile(good) as zin, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info.filename, zin.read(info))
        zout.writestr("ppt/media/bomb.bin", b"\0" * (48 * 1024 * 1024))
    (OUT / "hostile-zipbomb.pptx").write_bytes(bomb.getvalue())

    # Compound-file containers: an encrypted OOXML package and a legacy .ppt.
    def ole(stream_name: str) -> bytes:
        header = bytearray(512)
        header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
        header[0x1E:0x20] = (9).to_bytes(2, "little")
        header[0x30:0x34] = (0).to_bytes(4, "little")
        directory = bytearray(512)
        name = stream_name.encode("utf-16-le")
        directory[128 : 128 + len(name)] = name
        return bytes(header) + bytes(directory)

    (OUT / "encrypted.pptx").write_bytes(ole("EncryptedPackage"))
    (OUT / "legacy.ppt").write_bytes(ole("PowerPoint Document"))
    shutil.copy(good, OUT / "renamed-copy.ppsx")


def main() -> None:
    simple_text()
    images()
    shapes()
    tables()
    multi_slide()
    aspect_ratios()
    text_formatting()
    fills()
    groups()
    broken()
    print("fixtures written to", OUT)


if __name__ == "__main__":
    main()
